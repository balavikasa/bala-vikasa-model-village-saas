"""Workbook and CSV views over the exact, already-authorized report selection."""
from __future__ import annotations

import csv
from io import StringIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [('id', 'Village ID'), ('name', 'Village'), ('cluster', 'Cluster'), ('pc', 'PC'),
           ('assigned_da', 'Assigned DA'), ('status', 'Status'), ('visit_date', 'Visit date'),
           ('last_visit', 'Last visit'), ('submitted_by', 'Submitted by'), ('male', 'Male'),
           ('female', 'Female'), ('people_attended', 'Total attendances'), ('new_members', 'New members'),
           ('visits', 'Submissions'), ('report_url', 'Latest report')]


def safe_cell(value: Any) -> Any:
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def _sheet(wb: Workbook, name: str, columns: list[tuple[str, str]], rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(name)
    ws.append([label for _, label in columns])
    for row in rows:
        ws.append([safe_cell(row.get(key)) for key, _ in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='005C64')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for index, (key, _) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(index)].width = 26 if key in {'name', 'submitted_by', 'assigned_da', 'report_url'} else 19
    ws.row_dimensions[1].height = 28


def coverage_workbook(result: dict[str, Any]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Summary'
    ws.append(['Metric / filter', 'Value'])
    ws.append(['Period', result['period']['label']])
    ws.append(['Timezone', result['period']['timezone']])
    for prefix, values in [('Scope', result['counts']), ('Selection', result['selection_counts']), ('Filter', result['filters_applied'])]:
        for key, value in values.items():
            if key not in {'page', 'page_size'}:
                ws.append([f'{prefix}: {key}', safe_cell(value)])
    for key, value in result['definitions'].items():
        ws.append([key, value])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='005C64')
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 100
    ws.freeze_panes = 'A2'
    ws.sheet_view.showGridLines = False
    selected = result['all_items']
    _sheet(wb, 'Villages', COLUMNS, selected)
    # Only selected rows appear in these sheets, never rows hidden by the chosen
    # status, search or scope. Pagination alone does not constrain the export.
    for status in ('Done', 'Pending'):
        rows = [r for r in selected if r['status'] == status]
        if rows:
            _sheet(wb, status, COLUMNS, rows)
    _sheet(wb, 'Submissions', [('date', 'Entry date'), ('village', 'Village'), ('cluster', 'Cluster'),
           ('assigned_da', 'Assigned DA'), ('submitted_by', 'Submitted by'), ('male', 'Male'),
           ('female', 'Female'), ('people_attended', 'Total attendances'), ('new_members', 'New members'),
           ('other_village', 'Outside submitter assignment'), ('view_url', 'Report')], result['submissions'])
    # Team and monthly summaries describe the period/scope, independently of the
    # selected status. Separate sheet names make that boundary explicit.
    _sheet(wb, 'Team - scope', [('name', 'DA'), ('pc', 'PC'), ('cluster', 'Cluster'),
           ('assigned', 'Assigned villages'), ('done', 'Done by selected DA' if result['filters_applied'].get('submitted_da') else 'Done by any DA'), ('pending', 'Pending'),
           ('visits', 'Actual submissions'), ('male', 'Male'), ('female', 'Female'),
           ('people_attended', 'Total attendances'), ('new_members', 'New members'),
           ('other_visits', 'Outside own assignment')], result['team'])
    _sheet(wb, 'Monthly coverage - scope', [('month', 'Month'), ('villages', 'Active villages'),
           ('done', 'Done in month'), ('pending', 'Pending in month')], result['monthly_coverage'])
    return wb


def coverage_csv(result: dict[str, Any]) -> str:
    stream = StringIO(newline='')
    writer = csv.writer(stream)
    writer.writerow([label for _, label in COLUMNS])
    for row in result['all_items']:
        writer.writerow([safe_cell(row.get(key)) for key, _ in COLUMNS])
    return stream.getvalue()
