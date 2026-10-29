"""Authenticated reporting acceptance cases, executed by the normal Flask CI suite.

Local evidence runs separately: these route/auth tests require the project runtime.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.models import (
    DA,
    PC,
    AttendanceEntry,
    AttendanceStatus,
    Cluster,
    Committee,
    Role,
    User,
    Village,
    VillageAttendanceVisit,
)


@pytest.fixture()
def coverage_data(app, monkeypatch):
    monkeypatch.setattr('app.services.coverage.local_today', lambda: date(2026, 10, 29))
    monkeypatch.setattr('app.reports.local_today', lambda: date(2026, 10, 29))
    pc = PC(full_name='Coverage PC', cluster=Cluster.CSRB)
    other_pc = PC(full_name='Outside PC', cluster=Cluster.PDTC)
    da = DA(full_name='Coverage DA', pc=pc)
    other_da = DA(full_name='Zero-activity DA', pc=other_pc)
    admin = User(email='coverage-admin@example.invalid', display_name='Administrator', role=Role.ADMIN, password_hash='fixture')
    worker = User(email='coverage-da@example.invalid', display_name='Coverage DA', role=Role.DA, da=da, password_hash='fixture')
    coordinator = User(email='coverage-pc@example.invalid', display_name='Coverage PC', role=Role.PC, pc=pc, password_hash='fixture')
    db.session.add_all([pc, other_pc, da, other_da, admin, worker, coordinator])
    villages = []
    for i in range(43):
        village = Village(name=f'Village {i + 1:03}', da=da if i < 5 else other_da)
        committee = Committee(name='Water', committee_type='Water', village=village)
        db.session.add_all([village, committee])
        villages.append(village)
    db.session.add(Village(name='Disabled Village', da=da, is_enabled=False))
    db.session.flush()
    for i, (day, male, female, new_members) in enumerate([(28, 11, 13, 29), (29, 5, 5, 0)]):
        village = villages[i]
        parent = VillageAttendanceVisit(village=village, visit_date=date(2026, 10, day), visit_month=date(2026, 10, 1),
            new_members_count=new_members, submitted_by_user_id=worker.id,
            client_submission_id=f'coverage-parent-{i}', photo_path='fixture.webp', latitude=17.98, longitude=79.52)
        entry = AttendanceEntry(village=village, committee=village.committees[0], village_visit=parent,
            visit_date=parent.visit_date, male_count=male, female_count=female, status=AttendanceStatus.ON_TIME,
            submitted_by_user_id=worker.id, client_submission_id=f'coverage-child-{i}')
        db.session.add_all([parent, entry])
    db.session.commit()
    return {'admin': admin.id, 'da': worker.id, 'pc': coordinator.id, 'outside_da': other_da.id, 'da_profile': da.id}


def login_as(client, user_id):
    # Exercise Flask-Login's normal signed-session user loader, not a production bypass.
    with client.session_transaction() as session:
        session['_user_id'] = str(user_id)
        session['_fresh'] = True


def test_coverage_requires_authentication(client):
    assert client.get('/api/v1/coverage').status_code == 401
    assert client.get('/coverage/export.xlsx').status_code == 302


def test_full_pending_population_and_export_share_filters(client, coverage_data):
    login_as(client, coverage_data['admin'])
    response = client.get('/api/v1/coverage?period=this_month&status=pending&page=2')
    assert response.status_code == 200
    result = response.get_json()
    assert result['counts']['villages'] == 43
    assert result['counts']['villages_done'] == 2
    assert result['total'] == 41
    assert len(result['items']) == 16
    assert response.headers['Cache-Control'] == 'private, no-store'
    exported = client.get('/coverage/export.xlsx?period=this_month&status=pending&page=2')
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.data), data_only=True)
    assert workbook['Villages'].max_row == 42
    assert {r[5] for r in list(workbook['Villages'].values)[1:]} == {'Pending'}


def test_coverage_cannot_expand_da_or_pc_scope(client, coverage_data):
    login_as(client, coverage_data['da'])
    assert client.get('/api/v1/coverage').json['counts']['villages'] == 5
    assert client.get(f"/api/v1/coverage?da={coverage_data['outside_da']}").json['counts']['villages'] == 0
    login_as(client, coverage_data['pc'])
    assert client.get('/api/v1/coverage').json['counts']['villages'] == 5


def test_daily_counts_map_and_analytics_agree(client, coverage_data):
    login_as(client, coverage_data['admin'])
    for endpoint in ('/api/v1/coverage', '/api/v1/dashboard/series', '/api/v1/monitoring/visit-analytics'):
        response = client.get(endpoint + '?period=today')
        assert response.status_code == 200
        assert response.json['counts']['people_attended'] == 10
        assert response.json['counts']['villages_done'] == 1
    markers = client.get('/api/v1/monitoring/map?period=today').json['items']
    assert sum(row['status'] == 'Visited' for row in markers) == 1
    assert all(row['photo_url'] is None or '/photos/attendance/' in row['photo_url'] for row in markers)


def test_report_search_summary_and_export_are_same_population(client, coverage_data):
    login_as(client, coverage_data['admin'])
    query = '?period=this_month&type=Attendance&q=Village+001'
    response = client.get('/api/v1/reports' + query)
    assert response.status_code == 200
    assert len(response.json['items']) == 1
    assert response.json['summary']['people_attended'] == 24
    workbook = load_workbook(BytesIO(client.get('/reports/export.xlsx' + query).data), data_only=True)
    assert sum(row[12] or 0 for row in list(workbook['Reports'].values)[1:]) == 24
    assert sum(row[13] or 0 for row in list(workbook['Reports'].values)[1:]) == 29


def test_zero_activity_staff_and_invalid_period(client, coverage_data):
    login_as(client, coverage_data['admin'])
    result = client.get('/api/v1/coverage').json
    assert any(row['name'] == 'Zero-activity DA' and row['visits'] == 0 for row in result['team'])
    for endpoint in ('/api/v1/coverage', '/api/v1/reports', '/api/v1/monitoring/map', '/coverage/export.xlsx'):
        response = client.get(endpoint + '?period=custom&from=2026-10-29&to=2026-10-01')
        assert response.status_code == 422


def test_profile_invalid_dates_are_json_not_html(client, coverage_data):
    login_as(client, coverage_data['admin'])
    response = client.get(
        f"/directory/da/{coverage_data['da_profile']}",
        query_string={'period': 'custom', 'from': '<img src=x onerror=alert(1)>', 'to': '2026-10-29'},
    )
    assert response.status_code == 422
    assert response.is_json
    assert 'error' in response.json
